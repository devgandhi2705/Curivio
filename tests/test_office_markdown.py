"""
Word and PowerPoint uploads are stored as Markdown (backend/utils/office_markdown.py),
so chat previews can show them formatted.

The old docx extraction joined paragraphs with single newlines (one giant "paragraph",
hard-cut every 800 chars by the chunker) and appended every table at the END of the
document. PowerPoint wasn't accepted at all.
"""
import io

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from backend.utils.office_markdown import docx_to_markdown, pptx_to_markdown
from backend.services import document_extraction_service as chat_extract


def _docx() -> bytes:
    d = Document()
    d.add_heading("1. The project", 1)
    p = d.add_paragraph("We study ")
    p.add_run("road networks").bold = True
    p.add_run(" over time.")
    d.add_paragraph("First step", style="List Bullet")
    d.add_paragraph("Nested step", style="List Bullet 2")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Tool", "Role"
    t.cell(1, 0).text, t.cell(1, 1).text = "Spark", "compute | fast"
    d.add_heading("2. Theory", 2)
    d.add_paragraph("After the table.")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _pptx() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "OSM Road Networks"
    s.placeholders[1].text = "BDA Phase 1"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "The problem"
    tf = s.placeholders[1].text_frame
    tf.text = "Roads change over time"
    sub = tf.add_paragraph()
    sub.text, sub.level = "New segments", 1
    s.notes_slide.notes_text_frame.text = "Mention Pune\nand Mumbai"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Metrics"
    tb = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tb.cell(0, 0).text, tb.cell(0, 1).text = "Metric", "Meaning"
    tb.cell(1, 0).text, tb.cell(1, 1).text = "Degree", "Links per node"
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def test_word_keeps_headings_emphasis_lists_and_tables_in_document_order():
    md = docx_to_markdown(Document(io.BytesIO(_docx())))
    assert md.startswith("# 1. The project\n\nWe study **road networks** over time.")
    assert "- First step\n  - Nested step" in md
    assert "| Tool | Role |\n| --- | --- |\n| Spark | compute / fast |" in md
    # the table sits BEFORE the heading that follows it, not at the end
    assert md.index("| Tool |") < md.index("## 2. Theory") < md.index("After the table.")


def test_word_chunks_now_split_on_real_paragraphs():
    text, *_ = chat_extract.extract_document_text(_docx(), "guide.docx", ".docx")
    assert "\n\n" in text                                    # blocks, not one run-on line
    assert len([p for p in text.split("\n\n") if p.strip()]) >= 5


def test_powerpoint_becomes_one_section_per_slide_with_notes():
    md, structure = pptx_to_markdown(_pptx())
    assert md.startswith("## Slide 1: OSM Road Networks\n\n- BDA Phase 1")
    assert "## Slide 2: The problem\n\n- Roads change over time\n  - New segments" in md
    assert "> **Speaker notes:** Mention Pune and Mumbai" in md
    assert "## Slide 3: Metrics\n\n| Metric | Meaning |" in md
    assert [s["title"] for s in structure] == ["OSM Road Networks", "The problem", "Metrics"]


def test_powerpoint_is_accepted_by_chat():
    assert ".pptx" in chat_extract.DOCUMENT_EXTENSIONS
    text, pages, error, error_type = chat_extract.extract_document_text(_pptx(), "deck.pptx", ".pptx")
    assert error is None and "## Slide 2: The problem" in text and pages is None


def test_an_unreadable_or_empty_deck_says_so():
    _, _, error, error_type = chat_extract.extract_document_text(b"not a zip", "deck.pptx", ".pptx")
    assert error and error_type == chat_extract.ERROR_PPTX_UNREADABLE
    empty = io.BytesIO()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(empty)
