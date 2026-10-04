"""
Word (.docx) and PowerPoint (.pptx) files as Markdown for chat document previews.

Why Markdown rather than the flat text the extractors used to produce:
  * the preview can show the document FORMATTED (headings, lists, bold, tables,
    one card per slide) with the app's existing Markdown renderer, instead of a
    wall of text;
  * the old docx extraction joined paragraphs with single newlines and put every
    table at the END of the document. Chunking splits on blank lines, so a whole
    Word file was one "paragraph" hard-cut every 800 characters, mid-sentence,
    and tables were torn away from the text around them. Here blocks are kept in
    document order and separated by blank lines;
  * the models reading the chunks understand Markdown structure natively.

Pure functions over python-docx / python-pptx objects; no I/O beyond parsing the
bytes. Used by backend.services.document_extraction_service.
"""
from __future__ import annotations

import io
import re

_SLIDE_HEADING = "## Slide {n}"          # the frontend's slide view splits on this


# ── shared ────────────────────────────────────────────────────────────────────
def _cell(text: str) -> str:
    # A "|" inside a cell would split it (the app's renderer splits rows on "|").
    return re.sub(r"\s+", " ", (text or "").strip()).replace("|", "/") or " "


def _table_md(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head, body = rows[0], rows[1:]
    out = ["| " + " | ".join(_cell(c) for c in head) + " |",
           "| " + " | ".join("---" for _ in head) + " |"]
    out += ["| " + " | ".join(_cell(c) for c in r) + " |" for r in body]
    return "\n".join(out)


def _emphasis(text: str, bold: bool | None, italic: bool | None) -> str:
    if not text.strip() or not (bold or italic):
        return text
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    mark = "***" if bold and italic else "**" if bold else "*"
    return f"{lead}{mark}{core}{mark}{trail}"


def _tidy(s: str) -> str:
    # "**a****b**" (two adjacent bold runs) -> "**ab**"
    return s.replace("****", "").replace("******", "")


# ── Word ──────────────────────────────────────────────────────────────────────
def _docx_inline(paragraph) -> str:
    from docx.text.hyperlink import Hyperlink
    from docx.text.run import Run
    parts: list[str] = []
    try:
        items = list(paragraph.iter_inner_content())
    except AttributeError:          # older python-docx: runs only
        items = list(paragraph.runs)
    for item in items:
        if isinstance(item, Run):
            parts.append(_emphasis(item.text or "", item.bold, item.italic))
        elif isinstance(item, Hyperlink):
            text = item.text or ""
            url = getattr(item, "url", "") or ""
            parts.append(f"[{text}]({url})" if url and text.strip() else text)
    text = _tidy("".join(parts))
    return text if text.strip() else (paragraph.text or "")


def _docx_list_level(paragraph) -> int | None:
    """Indent level if this paragraph is a list item (numbering or a List style)."""
    ppr = paragraph._p.pPr
    num = ppr.numPr if ppr is not None else None
    if num is not None and num.numId is not None:
        ilvl = num.ilvl.val if num.ilvl is not None else 0
        return int(ilvl or 0)
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style.startswith("List"):
        digits = "".join(ch for ch in style if ch.isdigit())
        return max(0, int(digits) - 1) if digits else 0
    return None


def docx_to_markdown(document) -> str:
    """A python-docx Document as Markdown, blocks in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks: list[str] = []
    list_run: list[str] = []

    def flush_list():
        if list_run:
            blocks.append("\n".join(list_run))
            list_run.clear()

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            flush_list()
            table = Table(child, document)
            md = _table_md([[cell.text for cell in row.cells] for row in table.rows])
            if md:
                blocks.append(md)
            continue
        if tag != "p":
            continue
        p = Paragraph(child, document)
        plain = (p.text or "").strip()
        if not plain:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style == "Title":
            flush_list()
            blocks.append(f"# {plain}")
            continue
        if style.startswith("Heading"):
            flush_list()
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(6, int(digits)) if digits else 1
            blocks.append(f"{'#' * level} {plain}")
            continue
        level = _docx_list_level(p)
        if level is not None:
            marker = "1." if "Number" in style else "-"
            list_run.append(f"{'  ' * level}{marker} {_docx_inline(p).strip()}")
            continue
        flush_list()
        blocks.append(_docx_inline(p).strip())
    flush_list()
    return "\n\n".join(b for b in blocks if b.strip())


# ── PowerPoint ────────────────────────────────────────────────────────────────
def _shape_order(shape) -> tuple[int, int]:
    return (shape.top if shape.top is not None else 0, shape.left if shape.left is not None else 0)


def _pptx_shape_lines(shape, out: list[str]) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in sorted(shape.shapes, key=_shape_order):
            _pptx_shape_lines(sub, out)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        md = _table_md([[cell.text for cell in row.cells] for row in shape.table.rows])
        if md:
            out.append("")
            out.append(md)
            out.append("")
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = _tidy("".join(_emphasis(r.text or "", r.font.bold, r.font.italic) for r in para.runs)).strip()
            if text:
                out.append(f"{'  ' * (para.level or 0)}- {text}")


def pptx_to_markdown(file_bytes: bytes) -> tuple[str, list[dict]]:
    """A .pptx as Markdown — one "## Slide N: title" section per slide, its text as
    (nested) bullets, tables as tables, speaker notes as a quote — plus the slide
    titles as the document's structure [{title, level}]."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    sections: list[str] = []
    structure: list[dict] = []
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text or "").strip() if title_shape is not None and title_shape.has_text_frame else ""
        title = re.sub(r"\s+", " ", title)
        lines: list[str] = []
        for shape in sorted(slide.shapes, key=_shape_order):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            _pptx_shape_lines(shape, lines)
        if not title:                       # untitled slide: its first line names it
            first = next((l for l in lines if l.strip() and not l.startswith("|")), "")
            title = re.sub(r"^\s*-\s*", "", first)[:80]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        heading = _SLIDE_HEADING.format(n=n) + (f": {title}" if title else "")
        body = "\n".join(lines).strip()
        parts = [heading]
        if body:
            parts.append(body)
        if notes:
            parts.append("> **Speaker notes:** " + re.sub(r"\s*\n\s*", " ", notes))
        sections.append("\n\n".join(parts))
        structure.append({"title": title or f"Slide {n}", "level": 1})
    return "\n\n".join(sections), structure
